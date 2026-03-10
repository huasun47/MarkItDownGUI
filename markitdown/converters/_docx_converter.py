"""使用 python-docx 的 Word 文档转换器"""

import sys
import io
import base64
import zipfile
from typing import BinaryIO, Any
from warnings import warn

from ._html_converter import HtmlConverter
from .._base_converter import DocumentConverterResult
from .._stream_info import StreamInfo
from .._exceptions import MissingDependencyException, MISSING_DEPENDENCY_MESSAGE

# Try loading python-docx
_dependency_exc_info = None
try:
    from docx import Document
    from docx.shared import Inches
except ImportError:
    _dependency_exc_info = sys.exc_info()

ACCEPTED_MIME_TYPE_PREFIXES = [
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
]

ACCEPTED_FILE_EXTENSIONS = [".docx"]


def extract_images_from_docx(docx_stream: BinaryIO) -> list:
    """
    从 DOCX 文件中提取图片并转换为 base64
    
    Args:
        docx_stream: DOCX 文件流
        
    Returns:
        list: 包含 base64 图片数据的列表
    """
    images = []
    
    try:
        # 重新定位到文件开头
        docx_stream.seek(0)
        
        # 使用 zipfile 读取 DOCX 文件
        with zipfile.ZipFile(docx_stream, 'r') as zip_file:
            # 查找 media 文件夹中的图片
            for file_name in zip_file.namelist():
                if file_name.startswith('word/media/') and file_name.lower().endswith(('.png', '.jpg', '.jpeg', '.gif', '.bmp')):
                    # 提取图片
                    image_data = zip_file.read(file_name)
                    
                    # 获取文件扩展名确定 MIME 类型
                    ext = file_name.lower().split('.')[-1]
                    mime_type = {
                        'png': 'image/png',
                        'jpg': 'image/jpeg',
                        'jpeg': 'image/jpeg',
                        'gif': 'image/gif',
                        'bmp': 'image/bmp'
                    }.get(ext, 'image/png')
                    
                    # 转换为 base64
                    base64_data = base64.b64encode(image_data).decode('utf-8')
                    data_url = f"data:{mime_type};base64,{base64_data}"
                    images.append(data_url)
                    
    except Exception as e:
        warn(f"提取图片失败: {e}")
    
    return images


class DocxConverter(HtmlConverter):
    """
    使用 python-docx 转换 DOCX 文件到 Markdown，支持图片提取和 base64 转换。
    """

    def __init__(self):
        super().__init__()
        self._html_converter = HtmlConverter()

    def accepts(
        self,
        file_stream: BinaryIO,
        stream_info: StreamInfo,
        **kwargs: Any,  # Options to pass to the converter
    ) -> bool:
        mimetype = (stream_info.mimetype or "").lower()
        extension = (stream_info.extension or "").lower()

        if extension in ACCEPTED_FILE_EXTENSIONS:
            return True

        for prefix in ACCEPTED_MIME_TYPE_PREFIXES:
            if mimetype.startswith(prefix):
                return True

        return False

    def convert(
        self,
        file_stream: BinaryIO,
        stream_info: StreamInfo,
        **kwargs: Any,  # Options to pass to the converter
    ) -> DocumentConverterResult:
        # Check: the dependencies
        if _dependency_exc_info is not None:
            raise MissingDependencyException(
                MISSING_DEPENDENCY_MESSAGE.format(
                    converter=type(self).__name__,
                    extension=".docx",
                    feature="docx",
                )
            ) from _dependency_exc_info[
                1
            ].with_traceback(  # type: ignore[union-attr]
                _dependency_exc_info[2]
            )

        try:
            # 提取图片
            images = extract_images_from_docx(file_stream)
            
            # 重新定位到文件开头
            file_stream.seek(0)
            
            # 使用 python-docx 读取文档
            doc = Document(file_stream)
            
            # 转换为 HTML
            html_content = self._docx_to_html(doc, images)
            
            # 设置 keep_data_uris=True 以保留完整的 base64 图片数据
            kwargs['keep_data_uris'] = True
            
            return self._html_converter.convert_string(
                html_content,
                **kwargs,
            )
            
        except Exception as e:
            warn(f"DOCX 转换失败: {e}")
            # 回退到原始方法
            file_stream.seek(0)
            from ..converter_utils.docx.pre_process import pre_process_docx
            import mammoth
            
            pre_process_stream = pre_process_docx(file_stream)
            html_content = mammoth.convert_to_html(pre_process_stream).value
            
            return self._html_converter.convert_string(
                html_content,
                **kwargs,
            )

    def _docx_to_html(self, doc, images: list) -> str:
        """
        将 python-docx 文档对象转换为 HTML
        
        Args:
            doc: python-docx Document 对象
            images: 提取的图片列表
            
        Returns:
            str: HTML 内容
        """
        html_parts = ['<html><body>']
        image_index = 0
        
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                # 处理段落样式
                style = ""
                if paragraph.style.name.startswith('Heading'):
                    level = paragraph.style.name.replace('Heading ', '')
                    style = f" h{level}"
                    html_parts.append(f'<h{level}>{paragraph.text}</h{level}>')
                else:
                    html_parts.append(f'<p>{paragraph.text}</p>')
            
            # 检查是否包含图片
            for run in paragraph.runs:
                if hasattr(run.element, 'xpath') and run.element.xpath('.//pic:pic'):
                    if image_index < len(images):
                        html_parts.append(f'<img src="{images[image_index]}" alt="图片 {image_index + 1}" />')
                        image_index += 1
        
        # 处理表格
        for table in doc.tables:
            html_parts.append('<table border="1">')
            for row in table.rows:
                html_parts.append('<tr>')
                for cell in row.cells:
                    html_parts.append(f'<td>{cell.text}</td>')
                html_parts.append('</tr>')
            html_parts.append('</table>')
        
        html_parts.append('</body></html>')
        return ''.join(html_parts)
